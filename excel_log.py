"""Human-readable audit trail of apply attempts — one row per job applied
to (or attempted), meant for a person to open directly in Excel. Separate
from storage.py's SQLite wrapper on purpose: jobs.db is the machine-queried
source of truth (candidate selection, cap counting), this file is a
by-hand-readable export of the same events, richer in per-attempt detail
(the actual screening questions asked and answers given) that doesn't need
to be queryable and would be awkward to keep in relational form.
"""

import logging
from datetime import datetime

from openpyxl import Workbook, load_workbook

import config

log = logging.getLogger("excel_log")

EXCEL_PATH = str(config.BASE_DIR / "applications_log.xlsx")

_SHEET_NAME = "Applications"
_HEADERS = [
    "Timestamp",
    "Status",
    "Company",
    "Job Title",
    "Job ID",
    "URL",
    "Fit Score",
    "Fit Reason",
    "Outcome Reason",
    "Questions & Answers",
    "Dry Run",
    "External Apply URL",
]


def _load_or_create_workbook() -> Workbook:
    try:
        wb = load_workbook(EXCEL_PATH)
        ws = wb[_SHEET_NAME]
        # Migrate an existing file from before a header was added: append
        # any missing headers at the end so old rows/columns stay intact
        # and new columns just start empty for pre-existing rows.
        existing_headers = [c.value for c in next(ws.iter_rows(max_row=1))]
        for header in _HEADERS:
            if header not in existing_headers:
                ws.cell(row=1, column=ws.max_column + 1, value=header)
        return wb
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = _SHEET_NAME
        ws.append(_HEADERS)
        return wb


def _format_qa_log(qa_log: list[dict]) -> str:
    if not qa_log:
        return "(none)"
    lines = []
    for qa in qa_log:
        answer = qa["answer"] if qa["answer"] is not None else "SKIPPED (needs manual review)"
        line = f"Q: {qa['question']}\nA: {answer}"
        options = qa.get("options")
        if options:
            line += f"\n   (options offered: {', '.join(options)})"
        lines.append(line)
    return "\n\n".join(lines)


def log_application(
    *,
    company: str,
    title: str,
    job_id: str,
    url: str,
    fit_score,
    fit_reason: str,
    applied: bool,
    outcome_reason: str,
    qa_log: list[dict],
    dry_run: bool,
    external_url: str | None = None,
) -> None:
    """Appends one row per apply attempt (dry-run or real). Status is
    "Success" (applied=True), "Dry Run" (a simulated attempt — `dry_run` is
    True regardless of `applied`, which is always False for these), or
    "Failed" (a real attempt that didn't result in an application). Checked
    in that order so dry-run rows never show as "Failed" — reading 18
    dry-run previews as 18 real failures at a glance would be actively
    misleading. Everything else that distinguishes *why* it didn't apply
    (external-apply skip, needs manual review, already applied, etc.) lives
    in the Outcome Reason column, not a proliferation of status values."""
    wb = _load_or_create_workbook()
    ws = wb[_SHEET_NAME]

    if applied:
        status = "Success"
    elif dry_run:
        status = "Dry Run"
    else:
        status = "Failed"

    ws.append(
        [
            datetime.now().isoformat(timespec="seconds"),
            status,
            company,
            title,
            job_id,
            url,
            fit_score,
            fit_reason,
            outcome_reason,
            _format_qa_log(qa_log),
            "Yes" if dry_run else "No",
            external_url or "",
        ]
    )

    try:
        wb.save(EXCEL_PATH)
    except PermissionError:
        # Most likely cause: the file is open in Excel. Log and move on —
        # a missed audit row for one job should never stop the apply cycle.
        log.error("Couldn't write %s (open in another program?) — row for %s not saved.", EXCEL_PATH, job_id)
